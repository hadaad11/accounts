"""Rebuild the encrypted PAYLOAD in index.html from the Excel files on Google Drive.

usage: DASH_PASSWORD=... python build/build.py [--from DIR]
  --from DIR  read DIR/<tab id>.xlsx instead of downloading
"""
import ast
import base64
import datetime as dt
import hashlib
import io
import json
import os
import re
import secrets
import sys
import urllib.request
from pathlib import Path
from zoneinfo import ZoneInfo

from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

SOURCES = [
    {"id": "receivables", "title": "المستحقات والالتزامات", "kind": "receivables",
     "source": "حسابات العملاء.xlsx", "file_id": "1GXXThw4VUr5oaJElKduNFqIt91iqAUVM"},
    {"id": "work", "title": "أسعار الشغل", "kind": "generic",
     "source": "اسعار الشغل.xlsx", "file_id": "1Ejhvdfn7Zz7GJk2XSa8KKfU1Dto4Q8EE"},
    {"id": "suppliers", "title": "أسعار البنود للموردين", "kind": "generic",
     "source": "اسعار البنود للموردين.xlsx", "file_id": "16lRTkSEqoxU2I5VM75Obl5uuEDXHIX0X",
     "rename": {"Sheet1": "أسعار البنود"}},
]
HTML = Path(__file__).resolve().parent.parent / "index.html"
TZ = ZoneInfo("Africa/Cairo")
ITER = 250_000
PAYLOAD_RE = re.compile(r"const PAYLOAD = (\{.*?\});\n")
TOTAL_WORDS = {"total", "الإجمالي", "الاجمالي", "إجمالي", "اجمالي", "المجموع"}


def num(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# ------------------------------------------------------------ formulas
REF_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
SAFE_NODES = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Add, ast.Sub, ast.Mult, ast.Div,
              ast.USub, ast.UAdd, ast.Constant, ast.Call, ast.Name, ast.Load)


def formula_ast(f):
    def sub(m):
        r1, c1 = int(m[2]), column_index_from_string(m[1])
        if m[3]:
            return f"_rng({r1},{c1},{int(m[4])},{column_index_from_string(m[3])})"
        return f"_c({r1},{c1})"
    try:
        tree = ast.parse(REF_RE.sub(sub, f.lstrip("=")), mode="eval")
    except SyntaxError:
        return None
    for n in ast.walk(tree):
        if not isinstance(n, SAFE_NODES) or (isinstance(n, ast.Constant) and not is_num(n.value)):
            return None
        if isinstance(n, ast.Call) and not (isinstance(n.func, ast.Name) and n.func.id in ("_c", "_rng", "SUM")):
            return None
    return tree


def signed_terms(node, sign=1):
    """Flatten a +/- expression into [(sign, node)]."""
    if isinstance(node, ast.Expression):
        return signed_terms(node.body, sign)
    if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub)):
        return signed_terms(node.left, sign) + signed_terms(node.right, -sign if isinstance(node.op, ast.Sub) else sign)
    return [(sign, node)]


def const_terms(f):
    """'=80000+15000+...' -> [80000, 15000, ...]; None if it isn't a plain sum of numbers."""
    tree = f and formula_ast(f)
    if not tree:
        return None
    terms = signed_terms(tree)
    if all(isinstance(n, ast.Constant) for _, n in terms):
        return [s * n.value for s, n in terms]
    return None


# ------------------------------------------------------------ sheet model
class Sheet:
    def __init__(self, wf, wv):
        self.name = wf.title.strip()
        self.cells = {}
        for row in wf.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                f = c.value if isinstance(c.value, str) and c.value.startswith("=") else None
                v = wv[c.coordinate].value if f else c.value
                if isinstance(v, str):
                    v = v.strip()
                    if not v:
                        continue
                self.cells[(c.row, c.column)] = {"v": num(v), "f": f}
        for cell in self.cells.values():
            if cell["f"] and cell["v"] is None:
                cell["v"] = self.evaluate(cell["f"])

    def value(self, r, c):
        cell = self.cells.get((r, c))
        if not cell:
            return 0
        if cell["v"] is None and cell["f"]:
            cell["v"] = self.evaluate(cell["f"])
        return cell["v"] if is_num(cell["v"]) else 0

    def evaluate(self, f):
        tree = formula_ast(f)
        if not tree:
            return None
        flat = lambda a: [y for x in a for y in (x if isinstance(x, list) else [x])]
        env = {"_c": self.value, "SUM": lambda *a: sum(flat(a)),
               "_rng": lambda r1, c1, r2, c2: [self.value(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]}
        try:
            return num(eval(compile(tree, "<f>", "eval"), {"__builtins__": {}}, env))
        except (ZeroDivisionError, TypeError, RecursionError):
            return None

    def row(self, r):
        return sorted((c, cell) for (rr, c), cell in self.cells.items() if rr == r)

    def blocks(self):
        rows = sorted({r for r, _ in self.cells})
        out = []
        for r in rows:
            if out and r == out[-1][-1] + 1:
                out[-1].append(r)
            else:
                out.append([r])
        return out

    def pairs(self):
        """Each row -> (label, number, cell): first text cell and last number cell."""
        for r in sorted({r for r, _ in self.cells}):
            cells = self.row(r)
            label = next((x["v"] for _, x in cells if isinstance(x["v"], str)), None)
            nums = [x for _, x in cells if is_num(x["v"])]
            if nums:
                yield label, nums[-1]["v"], nums[-1]


def load(data):
    wf = load_workbook(io.BytesIO(data))
    wv = load_workbook(io.BytesIO(data), data_only=True)
    return [Sheet(ws, wv[ws.title]) for ws in wf.worksheets]


# ------------------------------------------------------------ receivables workbook
def is_total(label):
    return isinstance(label, str) and label.strip().lower() in TOTAL_WORDS


def parse_company(sh):
    rows = [(None if not lbl or lbl.lower() == sh.name.lower() else lbl, v, cell) for lbl, v, cell in sh.pairs()]
    ti = next((i for i, (lbl, _, _) in enumerate(rows) if is_total(lbl)), None)
    seq = iter(range(1, 1000))
    label = lambda lbl: lbl or f"دفعة {next(seq)}"
    if ti is not None:
        return rows[ti][1], {
            "kind": "items",
            "items": [{"label": label(l), "value": v} for l, v, _ in rows[:ti] if v],
            "notes": [{"label": l, "value": v} for l, v, _ in rows[ti + 1:] if l],
        }
    if len(rows) == 3 and not any(l for l, _, _ in rows) and abs(rows[0][1] - rows[1][1] - rows[2][1]) < 0.5:
        (_, due, dc), (_, paid, pc), (_, rest, _) = rows
        due_items = const_terms(dc["f"]) or [due]
        paid_items = const_terms(pc["f"]) or [paid]
        return rest, {
            "kind": "contract", "items": [], "notes": [], "due": due, "paid": paid,
            "due_items": [{"label": f"بند {i}", "value": v} for i, v in enumerate(due_items, 1)],
            "paid_items": [{"label": f"دفعة {i}", "value": v} for i, v in enumerate(paid_items, 1)],
        }
    items = [{"label": label(l), "value": v} for l, v, _ in rows if v]
    return sum(i["value"] for i in items), {"kind": "items", "items": items, "notes": []}


def parse_receivables(sheets):
    people, companies, debts = [], [], []
    for sh in sheets:
        if re.search(r"دين|ديون|التزام", sh.name):
            debts += [{"name": l, "value": v} for l, v, _ in sh.pairs() if l and not is_total(l)]
        elif re.search(r"[أا]فراد", sh.name):
            people += [{"name": l, "type": "person", "value": v} for l, v, _ in sh.pairs() if l and not is_total(l)]
        elif sh.cells:
            value, detail = parse_company(sh)
            companies.append({"name": sh.name, "type": "company", "value": value, "detail": detail})
    rec = people + companies
    tr, td = sum(r["value"] for r in rec), sum(d["value"] for d in debts)
    return {"receivables": rec, "debts": debts, "total_receivable": tr, "total_debt": td, "net": tr - td}


# ------------------------------------------------------------ generic workbook
def as_calc(sh, blk):
    """Two rows: labels over numbers, one number is a formula over the others."""
    if len(blk) != 2:
        return None
    head, vals = dict(sh.row(blk[0])), dict(sh.row(blk[1]))
    numc = [c for c, x in vals.items() if is_num(x["v"])]
    fc = [c for c in numc if vals[c]["f"]]
    if len(fc) != 1 or len(numc) < 3 or not all(isinstance(head.get(c, {}).get("v"), str) for c in numc):
        return None
    fc = fc[0]
    body = vals[fc]["f"]
    refs = {(int(m[2]), column_index_from_string(m[1])) for m in REF_RE.finditer(body)}
    if not refs or any(r != blk[1] or c not in numc for r, c in refs):
        return None
    title = next((x["v"] for c, x in sorted({**vals, **head}.items()) if c not in numc and isinstance(x["v"], str)), None)
    parts = [c for c in numc if c != fc]
    cols = parts + [fc]
    return {
        "title": title, "columns": [head[c]["v"] for c in cols], "rows": [[{"n": vals[c]["v"]} for c in cols]],
        "kind": "calc", "op": "x" if "*" in body and "+" not in body else "+",
        "total": vals[fc]["v"], "total_label": head[fc]["v"],
        "parts": [{"label": head[c]["v"], "value": vals[c]["v"]} for c in parts],
        "_at": (blk[1], fc),
    }


def as_result(sh, blk, regions):
    """One row: a label and a formula over calc totals -> a KPI with a waterfall."""
    if len(blk) != 1:
        return None
    cells = [x for _, x in sh.row(blk[0])]
    texts = [x["v"] for x in cells if isinstance(x["v"], str)]
    forms = [x for x in cells if x["f"] and is_num(x["v"])]
    if len(texts) != 1 or len(forms) != 1 or len(cells) != 2:
        return None
    kpi = {"label": texts[0], "value": forms[0]["v"]}
    tree = formula_ast(forms[0]["f"])
    totals = {r["_at"]: i for i, r in enumerate(regions) if "_at" in r}
    terms = signed_terms(tree) if tree else []
    at = lambda n: (n.args[0].value, n.args[1].value) if isinstance(n, ast.Call) and n.func.id == "_c" else None
    if terms and all(at(n) in totals for _, n in terms):
        kpi["steps"] = [{"label": regions[totals[at(n)]]["title"] or regions[totals[at(n)]]["total_label"],
                         "region": totals[at(n)], "value": regions[totals[at(n)]]["total"], "sign": s} for s, n in terms]
    return kpi


def as_table(sh, blk, title):
    cols = sorted({c for (r, c) in sh.cells if r in blk})
    first = dict(sh.row(blk[0]))
    has_head = len(blk) > 1 and all(isinstance(x["v"], str) for x in first.values())
    body = blk[1:] if has_head else blk
    defaults = iter(["البند", "التفاصيل"])
    columns = [first[c]["v"] if has_head and c in first else next(defaults, f"عمود {i + 1}") for i, c in enumerate(cols)]
    rows = []
    for r in body:
        cells = dict(sh.row(r))
        rows.append([None if c not in cells else {"n": cells[c]["v"]} if is_num(cells[c]["v"]) else {"s": str(cells[c]["v"])}
                     for c in cols])
    emph = [i for i, c in enumerate(cols) if any(sh.cells.get((r, c), {}).get("f") for r in body)]
    reg = {"title": title, "columns": columns, "rows": rows, "kind": "table", "emph": emph}
    if emph:
        reg["totals"] = {str(i): sum(x[i]["n"] for x in rows if x[i] and "n" in x[i]) for i in emph}
    return reg


def parse_generic(sheets):
    out = []
    for sh in sheets:
        if not sh.cells:
            continue
        regions, kpis, title = [], [], None
        for blk in sh.blocks():
            calc = as_calc(sh, blk)
            if calc:
                regions.append(calc)
                continue
            kpi = as_result(sh, blk, regions)
            if kpi:
                kpis.append(kpi)
                continue
            if len(blk) == 1 and len(sh.row(blk[0])) == 1 and isinstance(sh.row(blk[0])[0][1]["v"], str):
                title = sh.row(blk[0])[0][1]["v"]
                continue
            regions.append(as_table(sh, blk, title))
            title = None
        if not kpis:
            for r in regions:
                if r["kind"] == "table" and r["emph"]:
                    i = r["emph"][-1]
                    kpis.append({"label": f"إجمالي {r['columns'][i]}", "value": r["totals"][str(i)]})
        for r in regions:
            r.pop("_at", None)
        out.append({"name": sh.name, "regions": regions, "kpis": kpis})
    return {"sheets": out}


# ------------------------------------------------------------ io
def fetch(file_id):
    urls = [f"https://drive.google.com/uc?export=download&id={file_id}",
            f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx"]
    for url in urls:
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            data = urllib.request.urlopen(req, timeout=60).read()
        except OSError as e:
            print(f"  {url}: {e}", file=sys.stderr)
            continue
        if data[:2] == b"PK":
            return data
    sys.exit(f"couldn't download {file_id} — is it shared as 'Anyone with the link can view'?")


def derive(pw, salt, iters):
    return hashlib.pbkdf2_hmac("sha256", pw.encode(), salt, iters, 32)


def decrypt(p, pw):
    key = derive(pw, base64.b64decode(p["salt"]), p["iter"])
    return json.loads(AESGCM(key).decrypt(base64.b64decode(p["iv"]), base64.b64decode(p["ct"]), None))


def encrypt(doc, pw):
    salt, iv = secrets.token_bytes(16), secrets.token_bytes(12)
    ct = AESGCM(derive(pw, salt, ITER)).encrypt(iv, json.dumps(doc, ensure_ascii=False).encode(), None)
    b = lambda x: base64.b64encode(x).decode()
    return {"enc": 1, "salt": b(salt), "iv": b(iv), "iter": ITER, "ct": b(ct)}


def main():
    pw = os.environ.get("DASH_PASSWORD") or sys.exit("DASH_PASSWORD is not set")
    local = Path(sys.argv[sys.argv.index("--from") + 1]) if "--from" in sys.argv else None
    html = HTML.read_text(encoding="utf-8")
    m = PAYLOAD_RE.search(html)
    old = json.loads(m.group(1))
    try:
        prev = {t["id"]: t for t in decrypt(old, pw)["tabs"]} if old.get("enc") else {}
    except Exception:
        sys.exit("DASH_PASSWORD doesn't open the current dashboard")

    now = dt.datetime.now(TZ).strftime("%Y-%m-%d %H:%M")
    tabs, changed = [], []
    for s in SOURCES:
        data = (local / f"{s['id']}.xlsx").read_bytes() if local else fetch(s["file_id"])
        sheets = load(data)
        for sh in sheets:
            sh.name = s.get("rename", {}).get(sh.name, sh.name)
        parsed = parse_receivables(sheets) if s["kind"] == "receivables" else parse_generic(sheets)
        p = prev.get(s["id"])
        same = p is not None and p["data"] == json.loads(json.dumps(parsed))
        if not same:
            changed.append(s["title"])
        tabs.append({"id": s["id"], "title": s["title"], "kind": s["kind"], "source": s["source"],
                     "modified": p["modified"] if same else now,
                     "edit_url": f"https://docs.google.com/spreadsheets/d/{s['file_id']}/edit",
                     "data": parsed})

    if not changed and [t["id"] for t in tabs] == list(prev):
        print("no changes")
        return
    payload = json.dumps(encrypt({"generated": now, "tabs": tabs}, pw), ensure_ascii=False)
    HTML.write_text(html[:m.start(1)] + payload + html[m.end(1):], encoding="utf-8")
    print(f"updated {now}: {', '.join(changed)}")
    if os.environ.get("GITHUB_OUTPUT"):
        with open(os.environ["GITHUB_OUTPUT"], "a") as f:
            f.write(f"stamp={now}\n")


if __name__ == "__main__":
    main()
